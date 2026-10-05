"""Use bundled document Python, not the training venv."""
from pathlib import Path
import json,re,argparse
from datetime import datetime,timezone
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'deliverables';OUT.mkdir(exist_ok=True)
def read(name):return json.loads((ROOT/'reports'/name).read_text(encoding='utf-8'))
def pct(x):return f'{x*100:.1f}%'
def inline(p,text):
    pattern=r'(\[([^\]]+)\]\(([^)]+)\)|`([^`]+)`|\*\*([^*]+)\*\*)'
    pos=0
    for m in re.finditer(pattern,text):
        p.add_run(text[pos:m.start()])
        if m.group(2):
            url=m.group(3);link=OxmlElement('w:hyperlink')
            rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);link.set(qn('r:id'),rid)
            r=OxmlElement('w:r');prop=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'1F4E79');prop.append(c);r.append(prop);t=OxmlElement('w:t');t.text=m.group(2);r.append(t);link.append(r);p._p.append(link)
        else:
            r=p.add_run(m.group(4) or m.group(5));r.bold=bool(m.group(5))
            if m.group(4):r.font.name='Consolas';r.font.size=Pt(10)
        pos=m.end()
    p.add_run(text[pos:])
def setup(title):
    d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=Inches(.7);s.bottom_margin=Inches(.7);s.left_margin=Inches(.75);s.right_margin=Inches(.75)
    d.core_properties.title=title;d.core_properties.author='项目技术交付';d.core_properties.created=datetime.now(timezone.utc);d.core_properties.modified=datetime.now(timezone.utc)
    for border in list(d.styles.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Header','Footer']:
        style=d.styles[name];style.font.name='Microsoft YaHei';style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    normal=d.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.15
    d.styles['Title'].font.size=Pt(24);d.styles['Title'].paragraph_format.space_after=Pt(14)
    for name,size in [('Heading 1',15),('Heading 2',12)]:
        d.styles[name].font.size=Pt(size);d.styles[name].paragraph_format.space_before=Pt(14);d.styles[name].paragraph_format.space_after=Pt(7)
    h=s.header.paragraphs[0];h.text='中文居家机器人社会世界模型 V2';h.runs[0].font.size=Pt(9)
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('2026年9月25日  |  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    for r in f.runs:r.font.size=Pt(9)
    d.add_paragraph(title,'Title');return d
def table(d,rows):
    t=d.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=([2.5,1.3,1.5,1.7] if len(rows[0])==4 else [7/len(rows[0])]*len(rows[0]))
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(c,value) in enumerate(zip(cells,row)):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.08
            if j and len(str(value))<20:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,str(value))
            for r in p.runs:r.font.size=Pt(10.5);r.bold=i==0;r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
            pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders);fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'203864' if i==0 else 'F3F6FA' if i%2==0 else 'FFFFFF');pr.append(fill)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
        if i==0:
            repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    d.add_paragraph()
def markdown_doc(text,title,path):
    d=setup(title);lines=text.splitlines();i=0;code=False
    while i<len(lines):
        line=lines[i]
        if line.startswith('# '):i+=1;continue
        if line.startswith('```'):code=not code;i+=1;continue
        if code:
            p=d.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1
            r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(10)
            if i+1<len(lines) and not lines[i+1].startswith('```'):p.paragraph_format.keep_with_next=True
            i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                vals=[x.strip() for x in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x) for x in vals):rows.append(vals)
                i+=1
            table(d,rows);continue
        if not line.strip():i+=1;continue
        if line.startswith('!['):
            m=re.match(r'!\[([^]]*)\]\(([^)]+)\)',line)
            if m:
                d.add_picture(str(Path(path).parent/m.group(2)),width=Inches(6.9));caption=d.add_paragraph(m.group(1));caption.paragraph_format.space_after=Pt(10)
                for r in caption.runs:r.font.size=Pt(9.5)
                i+=1;continue
        if line.startswith('## '):d.add_paragraph(line[3:].replace(' / ','与').replace('｜',''),'Heading 1')
        elif line.startswith('### '):d.add_paragraph(line[4:],'Heading 2')
        elif line.startswith('- '):inline(d.add_paragraph(style='List Bullet'),line[2:])
        else:
            p=d.add_paragraph();inline(p,line)
            if line.startswith('Linux x86'):p.paragraph_format.keep_together=True
        i+=1
    d.save(path)

def report_markdown():
    cfg=read('train_config.json');ev=read('evaluation.json');base=read('tfidf_baselines.json');gpu=read('runtime_cuda.json');cpu=read('runtime_cpu.json');export=read('cpu_export.json');cal=read('calibration.json');transition=read('transition_fit.json')
    manifest=json.loads((ROOT/'data/processed/manifest.json').read_text(encoding='utf-8'));selection=json.loads((ROOT/'models/social_zh/best/selection.json').read_text());events=[json.loads(x) for x in (ROOT/'reports/train_events.jsonl').read_text().splitlines()]
    done=next(x for x in reversed(events) if x['type']=='complete');peak=max(x.get('peak_vram_mib',0) for x in events)
    names={'emotion':'当前情绪','dialog_act':'当前对话行为','intent':'中文任务意图','boundary':'居家边界','policy':'下一步策略','next_emotion':'下一轮情绪','next_act':'下一轮对话行为','task_domain':'任务领域'}
    rows='\n'.join(f'| {names[k]} | {x["n"]} | {pct(x["accuracy"])} | {pct(x["macro_f1"])} |' for k,x in ev['heads'].items())
    baseline_rows='\n'.join(f'| {names[k]} | {pct(ev["heads"][k]["macro_f1"])} | {pct(base["heads"][k]["macro_f1"])} | {(ev["heads"][k]["macro_f1"]-base["heads"][k]["macro_f1"])*100:+.1f} 个百分点 |' for k in names)
    forecasts='\n'.join(f'| {names[k]} | {pct(ev["heads"][k]["accuracy"])} | {pct(v["train_majority"]["accuracy"])} | {pct(v["oracle_current_label_persistence"]["accuracy"])} |' for k,v in ev['forecast_baselines'].items())
    ablations='\n'.join(f'| {name} | {ev["ablations"][mode]["next_emotion"]["n"]} | {pct(ev["ablations"][mode]["next_emotion"]["macro_f1"])} | {pct(ev["ablations"][mode]["next_act"]["macro_f1"])} |' for mode,name in [('full','完整输入'),('no_history','移除历史'),('no_candidate','移除候选回应'),('shuffled_candidate','打乱候选回应')])
    failures='；'.join(f'{r["name"]}→{r["action"]}' for r in gpu['cases'] if not r['passed']) or '本次预先写定的案例均符合允许动作集合'
    quant_rows='\n'.join(f'| {names[k]} | {v["n"]} | {pct(v["label_agreement"])} | {pct(v["cpu_accuracy"])} |' for k,v in cpu.get('quantization_check',{}).items())
    count_rows='\n'.join(f'| {source} | {manifest["splits"]["train"]["sources"].get(source,0)} | {manifest["splits"]["validation"]["sources"].get(source,0)} | {manifest["splits"]["test"]["sources"].get(source,0)} |' for source in ['cped_state','cped_transition','massive_zh','crosswoz','home_synthetic'])
    return f'''# 中文居家机器人社会世界模型训练与交付报告

本报告面向项目负责人，说明本次在 RTX 4060 笔记本上完成的中文社会世界模型训练、真实测试结果和机器人接入范围。交付采用 Qwen3.5-2B 文本主干，包含当前状态估计、候选回应条件下的下一轮状态预测、个人边界与偏好记忆、HTTP 服务和控制器请求适配。

当前结论是：已经形成可在本机运行和接入控制器联调的社会交互组件；“真实家庭机器人已经完成产品验收”仍不属于本次证据范围。世界预测只覆盖下一轮对话状态，训练依据是公开中文人际对话，不能据此声称掌握长期人类心理、真实动作因果或通用物理世界规律。

必须明确的是，本轮当前情绪分类及下一轮预测的宏 F1 偏低，尚不足以支撑自主社会规划。已完成的成果是可复现训练、本机推理与控制器协议联调；可靠的真实家庭社会世界预测仍未达到。预测接口保持 advisory_only，只作实验分析，不能把链路跑通等同于预测能力达标。

## 交付成果与核心结果

正式训练使用 {cfg['unique_train']:,} 条去重样本，执行 {cfg['epochs']} 轮、{done['step']:,} 个优化步骤，训练进程用时约 {done['elapsed_seconds']/60:.1f} 分钟。最佳检查点来自第 {selection['epoch']} 轮，由验证集八个任务头的平均宏 F1 选择。PyTorch 峰值分配显存 {peak/1024:.2f} GiB，驱动和桌面图形的额外占用不计入该统计。

中文任务意图测试准确率为 {pct(ev['heads']['intent']['accuracy'])}，宏 F1 为 {pct(ev['heads']['intent']['macro_f1'])}；下一轮情绪预测宏 F1 为 {pct(ev['heads']['next_emotion']['macro_f1'])}，下一轮对话行为预测宏 F1 为 {pct(ev['heads']['next_act']['macro_f1'])}。这些成绩分别对应不同监督任务，不能平均成一个“机器人社会智能准确率”。

GPU 完整状态更新实测 P50 为 {gpu['latency']['p50_ms']:.1f} ms、P95 为 {gpu['latency']['p95_ms']:.1f} ms；CPU 混合精度量化后分别为 {cpu['latency']['p50_ms']:.1f} ms 和 {cpu['latency']['p95_ms']:.1f} ms。二者均包含分词、模型、策略和短期状态管理。复杂文本生成另计，不能把快速分类延迟当成生成整段回答的延迟。

## 模型结构与能力边界

输入为中文 ASR 文本、最多四轮历史，以及机器人已有感知模块提供的结构化字段。单个 Qwen3.5 文本网络共享八个任务头；状态任务预测当前标签，转移任务读取当前话语和候选回应，预测原说话人下一轮的情绪与对话行为分布。没有额外加载视觉、语音或嵌入模型。

社会状态中的情绪、对话行为、任务意图、边界候选和策略来自模型；用户明示的边界、距离偏好与记忆来自状态库和规则。信任值不凭空生成，默认未知。既定拒绝不会因为一轮模型推测而自动取消。短期历史与长期记忆按用户隔离，长期存储需明确身份和记忆同意。

模型返回的行为通过有限动作表转换为控制器请求。移动还需授权、有效期与避障检查；设备还需目标设备和意图确认。模拟控制器已验证请求和回执流程，但机器人芯片、操作系统与 SDK 尚未确定，未进行真实电机或家庭设备验证。

## 底座与软件选型

选择 Qwen3.5-2B，固定官方权重版本，不把发布日期直接等同于任务效果。先下载 0.8B 与 2B 并实测反向传播，2B 在长度 256、批量 8 下留有显存余量，因此选用容量更大的 2B；未开展两种底座的同预算正式训练对照，不宣称已证明选型最优。仅加载文本部分，共 {cfg['total_parameters']:,} 参数，可训练 {cfg['trainable_parameters']:,} 参数。

使用 Transformers 5.17.0、PEFT 0.21.0 和 Torch 2.6.0 CUDA 12.4。Torch 版本按现有驱动兼容性固定，没有为了追新改动显卡驱动。LoRA 秩为 8，alpha 为 16，作用于注意力和 Gated Delta 投影；八个线性任务头共同优化。BF16 冻结主干配合 FP32 适配参数，启用梯度检查点和长度分桶。

CPU 导出对线性层做动态 INT8 量化，词嵌入保留 BF16，状态与归一化算子使用 FP32。文件约 {export['bytes']/1024**3:.2f} GiB。它是经过本机运行验证的 PyTorch 文件，不是通用 ONNX、RKNN 或 BPU 文件。10 TOPS 无法单独推断可用延迟，目标板卡还需要算子支持、内存带宽和编译测试。

## 中文数据调研与最终划分

CPED 用于中文情绪、对话行为和观察到的下一轮转移；MASSIVE zh-CN 用于家庭设备等任务意图；CrossWOZ 用于结合历史的任务领域识别。居家拒绝、恢复对话、取消与礼貌表达另做合成补充，明确保留合成来源。

| 来源 | 训练 | 验证 | 测试 |
|---|---|---|---|
{count_rows}

总计训练 {manifest['splits']['train']['count']:,} 条、验证 {manifest['splits']['validation']['count']:,} 条、测试 {manifest['splits']['test']['count']:,} 条。保留官方集合归属后做固定抽样与完整输入去重；CPED 的电视剧来源在三个集合之间互斥。训练中的合成样本共曝光五次，属于重采样，独立样本数没有乘五。

预测样本来自 A 说话、B 回应、A 再说话的真实语料顺序。输入包括前两步，目标只取第三步标签，第三步文本不进入输入。CPED 是电视剧，MASSIVE 和 CrossWOZ 也不是真实家庭机器人数据，因此本次结果不能替代部署域评测。性别、年龄、人格和演员身份未用于训练推断。

许可方面，Qwen3.5、CPED 仓库和 CrossWOZ 仓库采用 Apache-2.0 标识；MASSIVE 为 CC BY 4.0，随项目保留许可和来源。EmotionTalk 的非商业限制、SocialDial 未明确的数据再利用许可，使它们未被纳入本次产品方向训练。CPED 涉及影视来源，仓库许可不等于已完成第三方版权链核验。

## 训练过程与复现证据

随机种子固定为 20260925，学习率 1.5e-4，AdamW、余弦衰减和 5% 预热，批量 8，最长 256 token。不同数据只监督其已有标签，其余任务损失屏蔽；类别权重按训练频数平方根倒数缩放并截断。每轮结束保存 last 和优化器状态；验证得分提高时保存 best。

![训练日志中的损失和显存记录  损失下降不等同于真实家庭任务验收通过](../reports/training_curves.png)

数据与模型版本、文件字节数和 SHA256 记录在来源清单。数据审计未发现跨集合完整输入重复或 CPED 电视剧交叉，但无法证明公开语料未曾进入底座预训练。完整环境锁定文件、下载、构建、训练、评估、导出脚本随交付保留。

温度缩放仅在验证集拟合，用于减少置信度失真。校准不改变 argmax 类别，也不能保证对方心理被正确理解。策略在不确定时澄清，明示边界优先于分类输出。

## 测试结果与基线比较

下表为未参与梯度训练的清洗后测试子集。宏 F1 对每个类别同等计权，适合观察少数情绪和行为类别；准确率更受常见类别影响。居家边界和策略两项属于合成措辞测试，不能称为真实家庭成功率。

| 任务 | 测试数 | 准确率 | 宏 F1 |
|---|---|---|---|
{rows}

当前情绪头的 13 类宏 F1 为 19.5%，是本次明显短板。另将同一组情绪概率按 CPED 原始标签对应关系合并为负向、中性、正向三组，准确率为 54.8%、宏 F1 为 49.5%。该结果只是概率分组，不是新训练的独立情感头；其中“惊讶”按 CPED 的负向约定分组，不能视为普遍心理规律。细节见 `reports/coarse_affect.json`。

为检验神经模型的额外收益，使用相同训练划分训练字符 TF-IDF 与平衡线性 SVM。它不是另一套大模型，也没有使用测试集拟合词表。

| 任务 | 本模型宏 F1 | TF-IDF 宏 F1 | 差值 |
|---|---|---|---|
{baseline_rows}

## 下一轮世界预测的证据

除了低成本文本基线，预测任务还比较训练多数类，以及直接保持真实当前标签的 oracle 基线。oracle 使用了实际系统通常无法直接获得的真实当前标签，因此是一个较强参照。结果不支持的结论不会被“已训练世界模型”这个名称掩盖。

| 预测目标 | 本模型准确率 | 多数类准确率 | 真实当前标签保持 |
|---|---|---|---|
{forecasts}

概率预测还应检查负对数似然 NLL，数值越低越好，而不能只看最可能类别是否命中。下一轮情绪的模型 NLL 为 {ev['heads']['next_emotion']['nll']:.3f}，只使用训练标签频率的无条件基线为 {ev['forecast_baselines']['next_emotion']['train_marginal']['nll']:.3f}；下一轮行为分别为 {ev['heads']['next_act']['nll']:.3f} 和 {ev['forecast_baselines']['next_act']['train_marginal']['nll']:.3f}。基线概率只用训练集计数，测试标签仅用于评分。

在预先固定的同一组 500 条测试转移样本上，分别移除历史、移除候选回应和打乱候选回应。这里的输入干预只检验模型对条件的依赖，不是现实中的随机行动实验。

| 输入条件 | 样本数 | 下一轮情绪宏 F1 | 下一轮行为宏 F1 |
|---|---|---|---|
{ablations}

完整 per-class 结果、家庭场景子集、NLL、校准误差与按对话聚类的 bootstrap 区间保存在 `reports/evaluation.json`。如果移除候选输入没有明显损失，应据实理解为当前模型尚未充分利用候选行为，不能从一次演示中推出因果规划能力。

## 工程验证和量化影响

接口规则与存储测试覆盖明确拒绝和恢复、多人会话隔离、无同意不落盘、删除记忆和提醒、低置信度语音、移动授权、避障门控、重复与过期请求、云端不发送身份和历史以及超时降级。真实模型案例通过 {gpu['passed']}/{gpu['total']}；案例结果为：{failures}。这些是开发验收案例，不是盲测的人类评价。

最终自动化测试为 33 项通过。另启动真实 GPU HTTP 服务并配置鉴权，完成 10 项接口检查，覆盖边界保持、恢复聊天保留距离限制、候选预测、转述处理、参数拒收、记忆删除、未授权访问拒收和确认灯具请求的模拟回执；全部通过，结果见 `reports/http_smoke.json`。测试服务结束后已停止，正式服务由启动脚本开启。

真实运行曾发现两项工程问题：转述他人拒绝被策略头当成本人拒绝，以及已确认的设备意图因神经误分类而无法执行。现已加上转述门控，并由可信上游确认的白名单意图覆盖神经猜测，同时保留原始分数和覆盖原因。修改前结果仍保留在 runtime_cuda_before_policy_fix.json 与 runtime_cpu_before_policy_fix.json；这是规则修复，不应计为神经模型精度提高。GPU 与 CPU 均验证了灯具请求到模拟回执 acknowledged 的链路。

CPU 与 GPU 在固定 256 条公开测试样本上比较，下面每项的数量取决于这些样本实际拥有的监督标签。量化一致率不是全测试集的精度证明。

| 任务 | 对比数 | CPU 与 GPU 标签一致率 | CPU 子集准确率 |
|---|---|---|---|
{quant_rows}

下一轮预测接口给出的两个候选回应分布差异已用真实权重检查，开发案例中情绪分布 L1 差异为 {gpu['candidate_distribution_l1']:.4f}。存在差异本身不等于预测准确；准确性以独立测试与消融表为准。

本次还使用 CPED 官方训练部分的 {transition['training_triplet_occurrences']:,} 个对话三元组出现次数，学习小型条件转移表。这与本次神经转移训练使用的是同一组有效三元组，不重复计为新增训练样本。它不包含新的文本主干，推理时复用同一个编码器估计双方状态，再融合神经预测与转移先验。情绪先验权重为 {transition['heads']['next_emotion']['alpha']:.2f}，行为先验权重为 {transition['heads']['next_act']['alpha']:.2f}，均按验证集宏 F1 选择，随后在验证集校准温度。CPU 一致率表对比的是原始任务头；最终候选预测另包含这一融合步骤。

## 使用与负责人汇报口径

在本机双击 V2 的 `start_gpu.cmd`，健康检查地址为 http://127.0.0.1:8766/health，交互式接口文档位于 /docs。`examples/robot_client.py` 演示中文交互到模拟控制器回执的完整调用。CPU 版本有独立启动入口。详细字段、记忆删除、提醒轮询、动作授权、联网兜底及复现步骤见配套使用指南。

可向负责人表述为“完成面向中文居家机器人的社会状态与单步对话转移模型训练，并打通本机部署和控制器协议联调”。不宜表述为“已完成中国家庭生产验收”“已掌握真实人的内心”或“已经验证 10 TOPS 板卡实时运行”。

下一阶段的必要验证是选定目标芯片与 SDK，采集经同意的真实居家机器人交互轨迹，由独立标注者评估边界理解、澄清行为和下一轮状态预测，验证跨家庭、口音、噪声与多人员场景，并做受控实机试运行。本次交付为这些工作提供可运行的起点和可重复的基线。

## 主要来源与文件

- [Qwen3.5-2B 官方模型](https://huggingface.co/Qwen/Qwen3.5-2B)
- [CPED 作者数据仓库](https://github.com/scutcyr/CPED)
- [MASSIVE 官方数据集](https://huggingface.co/datasets/AmazonScience/massive)
- [CrossWOZ 作者仓库](https://github.com/thu-coai/CrossWOZ)
- [EmotionTalk 许可说明](https://huggingface.co/datasets/BAAI/Emotiontalk/blob/main/README.md)
- [SocialDial 作者仓库](https://github.com/zhanhl316/SocialDial)

模型权重位于 `models/social_zh/best`，底座位于 `models/qwen35_2B`，CPU 文件为 `models/social_zh/cpu_int8.pt`。数据来源清单为 `data/provenance.json`，划分摘要为 `data/processed/manifest.json`，训练事件为 `reports/train_events.jsonl`，模型评价为 `reports/evaluation.json`，运行验证为 `reports/runtime_cuda.json` 与 `reports/runtime_cpu.json`。
'''

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--guide-only',action='store_true');p.add_argument('--report-only',action='store_true');a=p.parse_args()
    if not a.report_only:
        guide=(ROOT/'README.md').read_text(encoding='utf-8');markdown_doc(guide,'中文居家机器人社会世界模型使用指南',OUT/'中文居家社会世界模型使用指南.docx')
    if not a.guide_only:
        text=report_markdown();(OUT/'训练与交付报告.md').write_text(text,encoding='utf-8');markdown_doc(text,'中文居家机器人社会世界模型训练与交付报告',OUT/'中文居家社会世界模型训练与交付报告.docx')
    print('Documents built')
