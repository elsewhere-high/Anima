"""Generate manager report and operating guide only from completed evidence."""
import json,importlib.util,statistics
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'deliverables'
def read(path):return json.loads((ROOT/path).read_text(encoding='utf-8'))
def pct(x):return f'{x*100:.2f}%'

def report():
    state=read('reports/current_state_evaluation.json');chat=read('reports/dialogue_repair_evaluation.json');review=read('reports/dialogue_repair_manual_review.json');release=read('models/release.json')
    review['summary']=review['summary'].rstrip('。')
    s=read('runs/current_state_lora/selection.json');d=read('runs/dialogue_sft/selection.json');data=read('data/dialogue/manifest.json');runtime=read('reports/runtime_cuda.json')
    repair=read('runs/dialogue_repair_v2/selection.json');repair_data=read('data/dialogue_repair_v2/manifest.json')
    metric_rows=[]
    for source,description in [('cped_state','CPED'),('chinese_medd','中文 MEDD')]:
        for k,name in [('emotion','情绪'),('dialog_act','对话行为'),('sentiment','正负情感')]:
            item=state.get('metrics',{}).get(source,{}).get(k)
            if item:
                a=item['original'];b=item['trained'];metric_rows.append(f'| {description} {name}（{b["n"]}条） | {pct(a["accuracy"])} / {pct(a["macro_f1"])} | {pct(b["accuracy"])} / {pct(b["macro_f1"])} |')
    metrics='\n'.join(metric_rows) if metric_rows else '| 未产生验证合格状态候选 | 保留 V2 | 未新增固定测试 |'
    medd=state.get('metrics',{}).get('chinese_medd',{}).get('emotion',{})
    active=medd.get('active_six_class_macro_f1')
    active_note=(f'按 MEDD 实际出现的六类另算宏 F1，原版为 {pct(active["original"])}，候选为 {pct(active["trained"])}；此数值不能直接与 CPED 的 13 类宏 F1 比较。' if active else '')
    intervals=[]
    for k,name in [('emotion','情绪'),('dialog_act','对话行为'),('sentiment','正负情感')]:
        item=state.get('metrics',{}).get('cped_state',{}).get(k)
        if item:
            ci=item['accuracy_delta_95ci'];intervals.append(f'{name}准确率差值的对话分组自举 95% 区间为 [{ci[0]*100:+.2f}, {ci[1]*100:+.2f}] 个百分点。')
    source_names={'grounded_authored':'自编场景','oasst2_zh':'OASST2 中文','opens2s_zh':'OpenS2S 中文'}
    nll='\n'.join(f'| {source_names.get(source,source)}（{a["examples"]}条） | {a["token_nll"]:.4f} | {chat["likelihood"]["trained"][source]["token_nll"]:.4f} |' for source,a in chat.get('likelihood',{}).get('base',{}).items())
    adopted='、'.join({'emotion':'当前情绪','dialog_act':'对话行为','sentiment':'正负情感'}[k] for k in release['state']['accepted_heads']) if release['state'] else '没有新状态层通过，保留原 V2'
    dialogue_status='采用本次监督微调的回复适配层' if release['dialogue'] else '本次回复候选没有通过完整发布检查；运行时保留原底座并接入多轮上下文'
    cpu=read('reports/runtime_cpu.json') if (ROOT/'reports/runtime_cpu.json').exists() else None
    cpu_text=(f'CPU 运行时检查通过，总检查耗时 {cpu["seconds"]:.1f} 秒，但质量等价未通过认定。固定验证探针中，CPU INT8 与 GPU 的情绪标签一致率只有 72.92%（96 条）；对话行为为 84.38%（64 条），正负情感为 83.87%（93 条）。这些是一致率而非正确率。CPU 版仅作为实验兼容入口，不推荐作为本次质量验收配置；未另做 CPU 全量精度与置信度校准。RTX 4060 GPU 服务是推荐入口，报告的公开测试成绩均来自 GPU。' if cpu else 'CPU 量化代码／导出状态见专项记录；本汇报不声称尚未完成的 CPU 实测。')
    return f'''# 中文社会交互模型 V4 训练与交付报告

## 本阶段交付结论

按 2026 年 9 月 26 日的新要求，本阶段聚焦当前情绪识别、对话行为、正负情感与中文多轮回应。未来情绪／行为预测和长期推演退出本阶段发布范围，旧版实验保留。

实际采用的新状态层：{adopted}。对话：{dialogue_status}。一个 Qwen3.5-2B 底座承载原有任务理解／策略、当前状态与中文回复的任务适配层；没有同时加载多个语言模型。模型已经完成的测试以本报告表格和原始 JSON 为准，不把训练完成等同于所有能力合格。

交付定位是可在本机运行、可接通用 HTTP/JSON 的社会交互组件。它不等于已在家庭现场验收的完整机器人大脑；视觉、语音识别、运动导航和真实设备执行结果仍由外部系统提供。

## 实际训练过程

硬件为 RTX 4060 Laptop 8 GiB、16 GB 内存。底座固定为 Qwen/Qwen3.5-2B 的 15852e8c16360a2fea060d615a32b45270f8a8fc 版本，避免运行中随上游更新改变权重。

当前状态共 95,740 条不同训练样本：93,176 条 CPED 与 2,564 条新增中文情绪样本。外部情绪重复四次，每轮 103,432 次样本曝光；联合更新 LoRA、13 类情绪、19 类对话行为和 3 类情感层。实际完成 {s['actual_steps']:,} 次优化更新、{s['epochs_completed']} 个完整轮次，用时 {s['seconds']/60:.1f} 分钟，峰值分配显存 {s['peak_vram_mib']/1024:.2f} GiB。显存为 PyTorch 分配量，不是整机显卡总占用。

状态训练最大长度 256，batch 16；适配层学习率 1e-5，分类层 5e-5，权重衰减 0.03。类别不均衡使用温和类别权重；只根据验证集选择检查点，温度只用验证集拟合。第一步已有非零 LoRA 梯度证据，实际更新了适配参数，而非仅改提示词。

CPED 情绪训练中“中性”有 31,390 条，“感谢”只有 204 条；对话行为中“非观点陈述”40,598 条，“问候”120 条。因此同时报告宏 F1、逐类别结果和训练多数类参照，避免只看总体准确率而忽略稀少类别；这说明数据分布，并非已证明所有错误的原因。

回复训练初筛 1,721 条；进一步抽检发现人工高评分资料也有知识错误和身份污染，因此在第一个训练步骤前收紧为 {data['files']['train']['n']} 条：{data['files']['train']['sources']['opens2s_zh']} 条 OpenS2S 合成情绪支持文本、16 条逐条审查的 OASST2 样本、48 条项目自编家庭对话目标。每轮通过 1／3／8 次重复形成 {data['files']['train']['sources']['opens2s_zh']+16*3+48*8} 次曝光。自编内容明确标注合成训练资料，不充当家庭测试证据。

回复训练使用新的 r=8 LoRA，学习率 3e-5，微批次 2，累积 8 次，最大长度 512。只有最终助手回复与结束标记计算损失；系统、用户、历史、填充和隐藏推理全部屏蔽。实际完成 {d['actual_steps']} 次优化更新、{d['actual_micro_steps']} 个微批次、{d['epochs_completed']} 轮，用时 {d['seconds']/60:.1f} 分钟，峰值分配显存 {d['peak_vram_mib']/1024:.2f} GiB。逐段词表投影的损失和梯度已与完整交叉熵作数值对照。

第一轮对话候选虽然降低了参考回复损失，但人工审查发现无图像时虚构衣服外观、未回答前文明确的物品位置，因此拒绝直接发布。随后加强文字输入边界提示，并补充上下文事实、纠正、缺失视觉、设备歧义和只倾听等合成场景。修复训练共 {repair_data['files']['train']['n']} 条不同记录，包含 285 条新增自编目标和原 472 条训练资料；同一场景的改写放在同一数据集合，完整提示及合成组的跨集合交集均为零。

修复从第一轮选中的适配层继续训练，学习率 2e-5，最大长度 768，微批次 2、累积 8 次，三轮上限；原 OpenS2S／OASST2／家庭目标／新增目标分别重复 1／2／2／3 次，每轮 1,391 次曝光。实际完成 {repair['actual_steps']} 次更新、{repair['epochs_completed']} 轮，用时 {repair['seconds']/60:.1f} 分钟，峰值分配显存 {repair['peak_vram_mib']/1024:.2f} GiB。选中第 {repair['selected']['epoch'] if repair['selected'] else '无'} 轮；依据公开验证损失保留、合成验证损失和验证生成检查选择，没有根据最终测试挑选轮次。两个中止的预检查实验及原因也保留，未用于发布。

![两阶段回复验证损失；阶段间系统提示发生变化，只比较各阶段内部趋势。](dialogue_training_curves.png)

## 固定测试结果

下面每格为“准确率 / 宏 F1”。准确率表示标对的样本比例，宏 F1 对每个类别等权平均，能暴露稀少类别表现弱的问题。原版与新候选使用同次执行数据对照；候选未通过的任务不会替换原版。所有情绪分类仍在 13 类输出空间完成，没有把粗粒度情感当成细粒度情绪成绩。

| 固定测试任务 | 原版 | 本次候选 |
|---|---|---|
{metrics}

{('![固定测试：原版与本次候选；实际是否采用见交付结论。](fixed_test.png)' if metric_rows else '')}

CPED 状态测试 2,398 条来自电视剧话语，曾在此前实验中分析过；中文 MEDD 本次新留出 312 条，来源包含日常、电影及合成内容。MEDD 真实标签覆盖六类，表中宏 F1 仍按全部 13 个输出类别计算；另外的六类宏 F1 见原始评估文件，不能与 13 类数值混报。

{active_note}

{' '.join(intervals)}CPED 改变量较小，区间跨零的任务不能据此宣称准确率已有确定提高；中文 MEDD 的较大增益也不等于真实家庭成绩。

正负情感是独立的 3 类任务。原版由校准后的情绪概率按 CPED 分组求和，本次候选直接监督学习原始情感标签。该数据中的“惊讶”按 CPED 体系计为负向，不代表一切现实惊讶都是负面。

意图、边界、策略和领域沿用原版适配层及分类权重。运行时实测：切换状态和生成任务再返回后，原版所有分类 logits 逐位相等。此证据证明任务路由没有改坏参数，不是新增意图准确率提升声明。

## 中文回复评估

公开测试集 82 条，包含 73 条合成参考和 9 条 OASST2 人工来源参考；修复阶段另有 35 条按场景隔离的自编测试记录。公开测试已经在第一轮使用过，自编测试与训练共享模板风格，均不应称为独立家庭验收。下表是目标回复每 token 的负对数似然，越低表示更贴近参考措辞，不是用户满意度，也不能证明参考知识正确。底座与修复候选使用完全相同的加强版系统提示，避免把提示词变化计成权重提升。

| 测试来源 | 原底座 NLL | 微调候选 NLL |
|---|---|---|
{nll}

另有第一轮训练前固定的 24 道开发探针，覆盖情绪倾听、混合情绪、名字／时间纠正、上下文指代、少追问、换话题、缺失视觉、设备未执行及知识问答。修复训练是针对已观察到的失败设计，因此这些题是已知回归测试。人工审查结论：{review.get('summary','详见逐题评审记录')}。项目助手逐题评分为底座 {review['total_points']['base']}／48、修复候选 {review['total_points']['trained']}／48；审查并非盲评或外部用户评价，不换算成“对话准确率”。全部原始回答和逐题理由随交付保留。

生成文本不能授权硬件。常规聊天保留完整会话上下文；设备、低语音识别置信度、请求安静、紧急协助等事件优先经过原控制规则。无实际执行结果时，额外检查明显的“已经打开／发送”等完成声明。该检查只是一层工程约束，不能保证识别所有语义幻觉。

修复后的 35 条合成测试中，简单字符串检查通过 33 条；这不是语义准确率。其中一条“看窗外的花”的回答仍虚构外观，其他个别回答建议使用尚未提供的图像能力。运行接口因此加入明确的缺失图像处理：识别到要求实际看颜色／外观、且不是讨论已给文字描述时，直接说明本接口没有图像。另对丧失倾诉中主动建议用新伙伴替代的已知错误作回复约束。这些属于部署工程修正，原始回答仍保留原分数；不能据此声称模型已学会识图或所有安慰场景都正确。

## 部署与实际运行

GPU 集成检查通过，用时 {runtime['seconds']:.1f} 秒，峰值分配显存 {runtime['peak_vram_mib']/1024:.2f} GiB。检查包括真实权重重载、任务适配层恢复、中文多轮回应、安静边界、低 ASR 置信度、设备确认、无同意不持久化及删除记忆。设备验证只在模拟控制器上执行，没有操纵真实家庭硬件。

完整服务还重跑了 24 道开发问答。修正了旧路由对高置信“澄清”类别不调用生成的问题，使位置、时间和名字纠正能真正到达用户；未经授权的设备请求可以根据上下文提问，但 task.authorized 仍为 false。未经核实的 CALL_HUMAN 分类仅作为建议，不构成联系他人的授权。实际 HTTP 测试返回了“你说眼镜在卧室的书桌”，并通过连续会话、安静／恢复和设备授权检查。自动测试共 122 项通过。

仍有明确错误：搬家送餐问题有时被回答成先确认是否方便，玻璃杯遇热开裂仍给出不正确的物理解释，部分故事逻辑和生活建议较弱。完整服务逐题记录见 integrated_dialogue_review.json，不能把“接口检查通过”理解为每个回答都正确。

{cpu_text}

V4 使用独立端口 8767 和独立记忆数据库。根目录 start_v4_gpu.ps1 启动本版，原 V2 启动入口继续保留。Windows 是本机实测环境；Linux 提供同一服务入口与依赖版本记录，未在本机运行 Linux／ARM／BPU 编译验收。

## 负责人需要知道的边界

当前情绪是基于可见中文文本的分类假设，不能保证每次判断正确，更不能推断用户未表达的隐藏心理。关系熟悉度来自交互计数或明确输入，信任值仍为未知；记忆以用户确认身份和同意为条件，不把自动猜测当成个人事实。

本次确实下载、筛选、训练、固定评估并接通运行链路，但公开数据到中国真实家庭仍有分布差异。后续现场验收应重点统计错误安慰、错误指代、反复追问、拒绝未遵从和设备误操作请求；本报告不捏造未采集的家庭测试成绩。

## 证据与来源

关键证据：runs/current_state_lora、runs/dialogue_sft、runs/dialogue_repair_v2 的协议、逐步日志和选择结果；reports/current_state_evaluation.json、dialogue_repair_evaluation.json、dialogue_repair_manual_review.json、runtime_cuda.json、data_audit.json；models/release.json 记录实际启用的参数。第一轮被拒绝的回复模型、审查和中止实验记录保留，便于复核。

![实际训练轨迹；验证曲线用于选择检查点，不能当作固定测试成绩。](training_curves.png)

模型及研究：[Qwen3.5-2B](https://huggingface.co/Qwen/Qwen3.5-2B)、[CPED](https://github.com/scutcyr/CPED)、[OpenS2S 作者项目](https://github.com/CASIA-LM/OpenS2S)、[OpenS2S 论文](https://arxiv.org/abs/2507.05177)、[OASST2](https://huggingface.co/datasets/OpenAssistant/oasst2)、[中文 MEDD](https://huggingface.co/datasets/Johnson8187/Chinese_Multi-Emotion_Dialogue_Dataset)。这里只使用 OpenS2S 的筛选文本，不宣称复现其语音系统。逐文件版本与哈希、许可声明及改动见数据 manifest 和 THIRD_PARTY_NOTICES.md。
'''

def guide():
    return '''# 中文社会交互模型 V4 使用指南

## 启动

在项目根目录打开 PowerShell，运行以下命令。服务加载完成后监听本机 8767 端口。

```powershell
.\\start_v4_gpu.ps1
```

推荐使用 RTX 4060 GPU 入口。CPU 入口 start_v4_cpu.ps1 属于实验兼容版：虽然运行检查通过，96 条验证探针上的情绪标签与 GPU 一致率仅 72.92%，不能沿用 GPU 精度声明。它使用共享 INT8 底座与小型 FP32 适配层。不要同时开多个模型服务或与训练争抢显存。启动前必须存在 v4/models/release.json；完成训练但没有通过评估的参数不会自动发布。

```powershell
Invoke-RestMethod http://127.0.0.1:8767/health
```

health 返回一个物理底座、实际采用的新状态层、是否加载回复适配层，以及 forecast_enabled=false。V4 没有提供未来推演接口。

## 中文多轮对话

服务启动后，另开一个终端，可直接运行交互客户端：

```powershell
.\\.venv_zh\\Scripts\\python.exe v4/scripts/chat_client.py
```

输入 /quit 退出。客户端使用固定会话连续交谈，同时显示当前情绪假设；默认不保存长期记忆。

本次交付时已在后台启动 GPU 服务，可先查 /health，直接运行客户端即可。根目录 stop_v4.ps1 可停止本次后台服务；以后用启动脚本在前台运行时按 Ctrl+C 停止。先停掉已有服务，再换 GPU／CPU 入口，避免占用相同端口。

发送 POST /v1/step，Content-Type 为 application/json。固定 user_id 和 session_id 才会延续同一个会话；下例可保存为 UTF-8 JSON 文件，用 curl.exe 的 --data-binary @文件名 发送。

```json
{
  "observation": {
    "user_id": "demo",
    "session_id": "home_chat",
    "speech": "今天开会总被打断，我只想说说，先不用建议。"
  }
}
```

随后用相同会话发送“刚才让我不舒服的是什么？”。读取 response 作为对外回复，state.emotion 是 13 类情绪假设，state.dialog_act 是当前对话行为，state.sentiment_cped 是另一个 3 类情感任务，均有分布和置信度。dialogue 记录本轮是否实际调用生成、历史长度、结束原因与耗时；不要把置信度当成心理事实。

当前分类读取最近四个历史话轮，回复生成最多读取最近十个历史话轮，并按 token 上限移除更早上下文。会话闲置五分钟过期；长对话需要另行提供明确保存的记忆，不能无限记住所有内容。

## 用户偏好和记忆

“请先别打扰我”会进入安静状态，“现在可以聊天了”重新允许对话。结构化 preferred_style 支持 neutral、warm、gentle、brief、silent。距离偏好与聊天是否开启分别保留，不因同意聊天自动授权靠近。

持久化需要 identity_verified=true、memory_consent=true，以及非 anonymous 的 user_id；仅在上游确实完成身份确认和取得同意后填写。memory_note 是明确要保存的内容，服务不会从生成文本自动发明记忆。未同意时仅保存本次会话临时状态。

DELETE /v1/users/{user_id} 删除该用户持久化内容、提醒和会话。提醒查询为 GET /v1/users/{user_id}/reminders；确认处理为 POST /v1/users/{user_id}/reminders/{reminder_id}/ack。不要把生成回答中说“记住了”当作数据库写入证明，应检查 memory_persisted 与实际查询结果。

## 对接机器人控制器

上游语音识别输出 speech 和 asr_confidence；视觉／传感器可提供 person_present、gaze、distance、pose、movement 等结构化输入。当前情绪分类器仍是文本模型，不会直接读取摄像头或波形。未知感知保留 unknown，不编造观察。

设备动作要求明确设备、受支持的意图和授权。例如打开客厅灯时，还需要 target_device="living_light"、confirmed_task_intent="iot_hue_lighton"、task_execution_authorized=true。只有语义模型猜中“开灯”不等于允许执行。

检查 task.authorized、task.execution_status 和 motion.authorized。调用 v2/social_world_zh/adapter.py 的 command_from_decision 可得到请求格式，再由真实控制器校验有效期、重复请求、设备白名单及避障条件。当前服务自身不直接开灯、移动或打电话；实际控制器应返回成功／失败结果，再据此向用户确认。未经核实的 CALL_HUMAN 分类只是模型建议，human_assistance.authorized=false、human_request=false，不能据此自动联系别人。

## 网络与平台

默认仅监听 127.0.0.1。本机对接不需要云端账户或推理密钥。需要其他机器访问时先设置环境变量 SOCIAL_API_TOKEN，再运行 v4/start.ps1 -ListenAddress 0.0.0.0；调用方使用 Authorization: Bearer 对应令牌。用户 ID 是业务标识，不能替代部署层的账号权限控制。

迁移到另一台 Windows／Linux x86 机器时使用 Python 3.12，重新建立虚拟环境。requirements.lock.txt 是本机完整版本记录，torch==2.6.0+cu124 需要使用 PyTorch 对应的 CUDA 12.4 软件源；CPU 专用环境应安装相应 CPU 构建并重新验证。不能假设直接在默认软件源安装这一锁定文件就一定成功，也不能直接复制 Windows 虚拟环境到 Linux。

迁移文件需保留项目相对目录：v2/models/qwen35_2B、v2/models/social_zh/best、v2/social_world_zh、v2/reports/calibration.json，以及 v4/social_v4、v4/models、release.json 指向的两个 v4/runs 检查点。不要只复制 LoRA 文件，适配层本身不是完整底座。Linux 环境准备好后执行 sh v4/start.sh；本机尚未完成 Linux、ARM 或专用机器人芯片的编译验收。

## 排错与复核

显存不足时关闭其他训练／服务后重启。出现 calibration 或 fingerprint mismatch 时，说明权重与校准文件不一致，不要手改校验值；应重新评估对应版本。服务未给出生成回复时先看 dialogue.status、安静边界、ASR 置信度和控制器事件类型。

训练日志在 v4/reports/current_state_training.log 和 dialogue_training.log。数据、协议、固定测试、逐题回复及最终权重选择分别在 data、runs、reports、models/release.json。训练脚本拒绝覆盖已有实验；复现实验应在独立目录保留原始文件和同一固定划分。

CPU 量化与 GPU 的真实集成记录见 reports/runtime_cpu.json、runtime_cuda.json。完整成绩和限制见同目录训练报告；不能用合成开发题通过率代替真实中国家庭验收。
'''

def main():
    OUT.mkdir(exist_ok=True)
    module_path=ROOT.parent/'v2/scripts/build_documents.py';spec=importlib.util.spec_from_file_location('document_renderer',module_path);renderer=importlib.util.module_from_spec(spec);spec.loader.exec_module(renderer)
    original_setup=renderer.setup
    def setup(title):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
        d=original_setup(title)
        for section in d.sections:
            section.header.paragraphs[0].text='中文居家社会交互模型 V4'
            footer=section.footer.paragraphs[0];footer.clear();footer.add_run('2026年9月26日  |  ')
            field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
            for run in footer.runs:run.font.size=Pt(9)
        return d
    renderer.setup=setup
    for name,title,content in [('训练与交付报告','中文社会交互模型 V4 训练与交付报告',report()),('使用指南','中文社会交互模型 V4 使用指南',guide())]:
        (OUT/(name+'.md')).write_text(content,encoding='utf-8');renderer.markdown_doc(content,title,OUT/(name+'.docx'))
        from docx import Document
        from docx.oxml import OxmlElement
        doc=Document(OUT/(name+'.docx'))
        for table in doc.tables:
            for i,row in enumerate(table.rows):
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next=i<len(table.rows)-1
        doc.save(OUT/(name+'.docx'))
    print(str(OUT))

if __name__=='__main__':main()
